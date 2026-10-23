"""Independent geometry checks and a batch test for a nonstandard EEM grid."""
from __future__ import annotations

import math
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from eem_tool import calculate_fri, interpolate, main, read_eem


def exact_region_area(ex_min, ex_max, em_min, em_max):
    """Integrate the Em>Ex domain analytically in one rectangular region."""
    left, right = max(ex_min, 0), ex_max
    flat_right = min(right, em_min)
    area = max(0, flat_right - left) * (em_max - em_min)
    slope_left, slope_right = max(left, em_min), min(right, em_max)
    if slope_right > slope_left:
        area += em_max * (slope_right - slope_left) - (slope_right**2 - slope_left**2) / 2
    return area


def test():
    old_ex, old_em, old_x = read_eem(ROOT / "examples/expected/黑_插值后.xlsx")
    old_result, old_areas = calculate_fri(old_x, old_ex, old_em)
    assert old_areas == [4000, 2500, 11000, 8450, 47800]
    expected_p = [0.18153525188254935, 0.3286844561702139,
                  0.13662650346003324, 0.2872406087016183,
                  0.06591317978558511]
    assert all(math.isclose(a, b, rel_tol=0, abs_tol=1e-12)
               for a, b in zip(old_result[10:15], expected_p))

    ex = [205, 225, 245, 265, 290, 320, 365, 410, 480]
    em = [245, 260, 280, 305, 325, 340, 360, 385, 420, 460, 505, 560, 610]
    x = [[10 + e / 100 + m / 200 for e in ex] for m in em]
    x[0][2] = 9999.9  # first Rayleigh band
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        book = Workbook()
        sheet = book.active
        sheet.title = "EEM"  # first sheet fallback
        sheet.append([None] + list(reversed(ex)))
        for m, row in zip(reversed(em), reversed(x)):
            sheet.append([m] + list(reversed(row)))
        source = folder / "different_grid.xlsx"
        book.save(source)
        loaded_ex, loaded_em, loaded_x = read_eem(source)
        assert loaded_ex == ex and loaded_em == em
        assert all(math.isclose(loaded_x[i][j], x[i][j], rel_tol=0, abs_tol=1e-12)
                   for i in range(len(em)) for j in range(len(ex)))
        assert main(folder) == 0
        out = folder / "EEM_处理结果"
        processed_ex, processed_em, processed = read_eem(out / "01_插值矩阵" / source.name)
        assert processed_ex == ex and processed_em == em
        assert processed[0][2] != 9999.9
        _, areas = calculate_fri(processed, ex, em)
        expected = [
            exact_region_area(ex[0], 250, em[0], 330),
            exact_region_area(ex[0], 250, 330, 380),
            exact_region_area(ex[0], 250, 380, em[-1]),
            exact_region_area(250, ex[-1], em[0], 380),
            exact_region_area(250, ex[-1], 380, em[-1]),
        ]
        assert all(math.isclose(a, b, rel_tol=0, abs_tol=1e-8)
                   for a, b in zip(areas, expected)), (areas, expected)
        packaged = os.environ.get("EEM_EXE")
        if packaged:
            subprocess.run([packaged, str(folder)], check=True)
            exe_ex, exe_em, exe_processed = read_eem(out / "01_插值矩阵" / source.name)
            assert exe_ex == ex and exe_em == em
            assert exe_processed == processed
        book = load_workbook(out / "02_五区积分汇总.xlsx", read_only=True, data_only=True)
        try:
            assert book["积分结果"].max_row == 2
            assert book["分区面积"].max_row == 6
            assert all(math.isclose(book["分区面积"].cell(i + 2, 3).value, area,
                                    rel_tol=0, abs_tol=1e-8)
                       for i, area in enumerate(expected))
        finally:
            book.close()
    print("variable-grid test passed; areas:", areas)

    dense_ex = list(range(200, 501, 5))
    dense_em = list(range(250, 601, 5))
    dense_x = [[1.0 + math.exp(-((e - 280) / 35) ** 2 - ((m - 340) / 55) ** 2)
                for e in dense_ex] for m in dense_em]
    dense_corrected, _ = interpolate(dense_x, dense_ex, dense_em)
    dense_result, dense_areas = calculate_fri(dense_corrected, dense_ex, dense_em)
    assert len(dense_corrected) == 71 and len(dense_corrected[0]) == 61
    assert math.isclose(sum(dense_areas), 73750, rel_tol=0, abs_tol=1e-8)
    assert math.isclose(sum(dense_result[10:15]), 1.0, rel_tol=0, abs_tol=1e-12)
    print("dense-grid test passed")


if __name__ == "__main__":
    test()
