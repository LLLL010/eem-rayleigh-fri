"""CSV orientation and end-to-end regression checks; uses no private sample files."""
from __future__ import annotations

import csv
import math
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from eem_tool import interpolate, main, read_eem
from eem_metrics import METRIC_HEADERS, calculate_metrics


def write_csv(path, header, axis, matrix, *, delimiter=",", preamble=None):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        if preamble:
            stream.write(preamble + "\n")
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow([""] + header)
        for wavelength, values in zip(axis, matrix):
            writer.writerow([wavelength] + values)


def close_matrix(left, right):
    return len(left) == len(right) and all(len(a) == len(b) for a, b in zip(left, right)) and all(
        math.isclose(a, b, rel_tol=0, abs_tol=1e-9)
               for rows in zip(left, right) for a, b in zip(*rows))


def test():
    ex, em, source = read_eem(ROOT / "examples/input/示例数据.xlsx")
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        standard = folder / "standard.csv"
        shimadzu = folder / "shimadzu.csv"
        picoquant = folder / "picoquant.csv"
        write_csv(standard, ex, em, source)
        # Shimadzu: Ex rows, Em columns, plus one 0 nm placeholder column.
        write_csv(shimadzu, [0] + em, ex,
                  [[0] + [source[i][j] for i in range(len(em))]
                   for j in range(len(ex))])
        # PicoQuant: one description line, wavelengths in metres, Ex rows/Em columns.
        write_csv(picoquant, [f"{v * 1e-9:.12g}" for v in em],
                  [f"{v * 1e-9:.12g}" for v in ex],
                  [[source[i][j] for i in range(len(em))] for j in range(len(ex))],
                  delimiter="\t", preamble="PicoQuant F300")
        for path in (standard, shimadzu, picoquant):
            actual_ex, actual_em, actual = read_eem(path)
            assert close_matrix([actual_ex], [ex]), path
            assert close_matrix([actual_em], [em]), path
            assert close_matrix(actual, source), path
        limited_ex = [250, 280, 310, 370, 450]
        limited_em = [380, 400, 430, 470, 520, 550]
        limited_x = [[1 + excitation / 200 + wavelength / 100
                      for excitation in limited_ex] for wavelength in limited_em]
        limited_shimadzu = folder / "shimadzu-limited.csv"
        write_csv(limited_shimadzu, limited_em, limited_ex,
                  [[limited_x[i][j] for i in range(len(limited_em))]
                   for j in range(len(limited_ex))])
        actual_ex, actual_em, actual = read_eem(limited_shimadzu)
        assert actual_ex == limited_ex and actual_em == limited_em
        assert close_matrix(actual, limited_x)
        assert main(folder) == 0
        result = folder / "EEM_处理结果"
        for path in (standard, shimadzu, picoquant, limited_shimadzu):
            assert (result / "01_插值矩阵" / (path.name + ".xlsx")).exists()
        book = load_workbook(result / "02_五区积分汇总.xlsx", read_only=True, data_only=True)
        try:
            rows = list(book["积分结果"].iter_rows(min_row=2, values_only=True))
            assert len(rows) == 4
            for row in rows:
                if row[0] == limited_shimadzu.name:
                    assert all(value is None for value in row[1:])
                    continue
                assert all(math.isclose(a, b, rel_tol=0, abs_tol=1e-9)
                           for a, b in zip(row[11:16],
                                           [0.18153525188254935, 0.3286844561702139,
                                            0.13662650346003324, 0.2872406087016183,
                                            0.06591317978558511]))
            metric_rows = list(book["荧光指标"].values)
            assert metric_rows[0] == tuple(METRIC_HEADERS)
            assert len(metric_rows) == 5
            corrected, _ = interpolate(source, ex, em)
            expected_metrics = calculate_metrics(corrected, ex, em)
            for row in metric_rows[1:]:
                if row[0] == limited_shimadzu.name:
                    continue
                assert row[0] in {path.name for path in (standard, shimadzu, picoquant)}
                for key, actual in zip(METRIC_HEADERS[1:], row[1:]):
                    expected = expected_metrics[key]
                    if isinstance(expected, float):
                        assert math.isclose(actual, expected, abs_tol=1e-9)
                    else:
                        assert actual == expected or (actual is None and expected == "")
        finally:
            book.close()
        print("CSV standard, Shimadzu and PicoQuant inputs passed")


if __name__ == "__main__":
    test()
