"""Hitachi Data Points TXT parsing and full batch processing on a nonstandard grid."""
from __future__ import annotations

import math
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from eem_tool import main, read_eem


def test():
    ex = [205, 225, 245, 265, 290, 320, 365, 410, 480]
    em = [245, 260, 280, 305, 325, 340, 360, 385, 420, 460, 505, 560, 610]
    values = [[10 + e / 100 + m / 200 for e in ex] for m in em]
    rows = [
        "Sample:\t'Synthetic",
        "Model:\tF-4600 FL Spectrophotometer",
        "EX Start WL:\t205.0 nm",
        "EM Start WL:\t245.0 nm",
        "Peaks",
        "1\t265.0/340.0\t999.0",
        "",
        "Data Points",
        "\t" + "\t".join(map(str, ex)),
    ]
    rows += ["\t".join([str(wavelength)] + [str(value) for value in row])
             for wavelength, row in zip(em, values)]
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        source = folder / "synthetic(FD3).TXT"
        source.write_text("\n".join(rows), encoding="utf-8")
        actual_ex, actual_em, actual = read_eem(source)
        assert actual_ex == ex and actual_em == em
        assert all(math.isclose(a, b, rel_tol=0, abs_tol=1e-12)
                   for source_row, actual_row in zip(values, actual)
                   for a, b in zip(source_row, actual_row))
        assert main(folder) == 0
        output = folder / "EEM_处理结果"
        assert (output / "01_插值矩阵" / "synthetic(FD3).TXT.xlsx").exists()
        book = load_workbook(output / "02_五区积分汇总.xlsx", read_only=True, data_only=True)
        try:
            assert book["积分结果"].max_row == 2
            assert book["荧光指标"].max_row == 2
            assert book["运行记录"].cell(2, 2).value == "成功"
        finally:
            book.close()
    print("Hitachi TXT nonstandard-grid batch test passed")


if __name__ == "__main__":
    test()
