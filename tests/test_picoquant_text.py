"""PicoQuant metre-unit text exports with no blank header cell."""
from __future__ import annotations

import csv
import math
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from eem_tool import main, read_eem


def write_picoquant(path, delimiter, ex, em, matrix):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        stream.write("PicoQuant F300, export\n")
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow([f"{v * 1e-9:.12g}" for v in em])
        for j, wavelength in enumerate(ex):
            writer.writerow([f"{wavelength * 1e-9:.12g}"] +
                            [matrix[i][j] for i in range(len(em))])


def test():
    ex, em, source = read_eem(ROOT / "examples/input/示例数据.xlsx")
    with tempfile.TemporaryDirectory() as directory:
        folder = Path(directory)
        files = [folder / "picoquant-tab.txt", folder / "picoquant-comma.dat"]
        write_picoquant(files[0], "\t", ex, em, source)
        write_picoquant(files[1], ",", ex, em, source)
        for path in files:
            actual_ex, actual_em, actual = read_eem(path)
            assert actual_ex == ex and actual_em == em
            assert all(math.isclose(a, b, rel_tol=0, abs_tol=1e-9)
                       for source_row, actual_row in zip(source, actual)
                       for a, b in zip(source_row, actual_row))
        limited_ex = [250, 280, 310, 370, 450]
        limited_em = [380, 400, 430, 470, 520, 550]
        limited_x = [[1 + wavelength / 100 + excitation / 200
                      for excitation in limited_ex] for wavelength in limited_em]
        limited_file = folder / "picoquant-limited.dat"
        write_picoquant(limited_file, "\t", limited_ex, limited_em, limited_x)
        assert main(folder) == 0
        result = folder / "EEM_处理结果"
        for path in files + [limited_file]:
            assert (result / "01_插值矩阵" / (path.name + ".xlsx")).exists()
        book = load_workbook(result / "02_五区积分汇总.xlsx", read_only=True, data_only=True)
        try:
            assert book["积分结果"].max_row == 4
            assert book["荧光指标"].max_row == 4
            integral_rows = {row[0]: row for row in list(book["积分结果"].values)[1:]}
            metric_rows = {row[0]: row for row in list(book["荧光指标"].values)[1:]}
            log_rows = {row[0]: row for row in list(book["运行记录"].values)[1:]}
            assert all(value is None for value in integral_rows[limited_file.name][1:])
            assert metric_rows[limited_file.name][1] is not None
            assert log_rows[limited_file.name][1] == "部分完成"
        finally:
            book.close()
    print("PicoQuant comma-DAT, tab-TXT and limited-range tests passed")


if __name__ == "__main__":
    test()
