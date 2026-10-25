"""Batch Rayleigh interpolation and five-region EEM integration."""
from __future__ import annotations

import bisect
import math
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from eem_readers import read_csv_eem, read_hitachi_txt
from eem_metrics import METRIC_HEADERS, calculate_metrics


EX = list(range(200, 501, 10))
EM = list(range(250, 601, 2))
AREAS = [4000, 2500, 11000, 8450, 47800]
MF = [sum(AREAS) / area for area in AREAS]
REGIONS = [
    ([j for j, ex in enumerate(EX) if ex <= 250], [i for i, em in enumerate(EM) if em <= 330]),
    ([j for j, ex in enumerate(EX) if ex <= 250], [i for i, em in enumerate(EM) if 330 <= em <= 380]),
    ([j for j, ex in enumerate(EX) if ex <= 250], [i for i, em in enumerate(EM) if em >= 380]),
    ([j for j, ex in enumerate(EX) if ex >= 250], [i for i, em in enumerate(EM) if em <= 380]),
    ([j for j, ex in enumerate(EX) if ex >= 250], [i for i, em in enumerate(EM) if em >= 380]),
]


def read_xlsx_eem(path: Path) -> tuple[list[float], list[float], list[list[float]]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Sheet1"] if "Sheet1" in book.sheetnames else book.worksheets[0]
        if sheet.max_row < 3 or sheet.max_column < 3:
            raise ValueError("至少需要两个 Ex 和两个 Em 波长点")
        if sheet.max_row * sheet.max_column > 1_000_000:
            raise ValueError("工作表过大；请移除数据区外的格式或内容")
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if rows[0][0] not in (None, ""):
        raise ValueError("A1 应为空白")
    def numeric(value, position):
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"{position} 应为有限数值；公式请先保存计算结果")
        return float(value)
    ex = [numeric(v, f"第一行第 {j+2} 列") for j, v in enumerate(rows[0][1:])]
    em = [numeric(row[0], f"第一列第 {i+2} 行") for i, row in enumerate(rows[1:])]
    x = [[numeric(v, f"第 {i+2} 行第 {j+2} 列") for j, v in enumerate(row[1:])]
         for i, row in enumerate(rows[1:])]
    return ex, em, x


def read_eem(path: Path) -> tuple[list[float], list[float], list[list[float]]]:
    """Read a standard workbook, supported CSV, or Hitachi text export."""
    if path.suffix.lower() == ".xlsx":
        ex, em, x = read_xlsx_eem(path)
    elif path.suffix.lower() == ".csv":
        ex, em, x = read_csv_eem(path)
    elif path.suffix.lower() == ".txt":
        ex, em, x = read_hitachi_txt(path)
    else:
        raise ValueError("支持的输入后缀为 .xlsx、.csv、.txt")
    if len(ex) < 2 or len(em) < 2 or len(x) != len(em) or any(len(row) != len(ex) for row in x):
        raise ValueError("强度矩阵须为至少 2×2，行数对应 Em、列数对应 Ex")
    if any(not math.isfinite(v) for v in ex + em) or any(
            not math.isfinite(v) for row in x for v in row):
        raise ValueError("波长或强度矩阵含非有限数值")
    if all(ex[i] > ex[i+1] for i in range(len(ex)-1)):
        ex.reverse()
        x = [row[::-1] for row in x]
    if all(em[i] > em[i+1] for i in range(len(em)-1)):
        em.reverse()
        x.reverse()
    if not all(ex[i] < ex[i+1] for i in range(len(ex)-1)) or not all(
            em[i] < em[i+1] for i in range(len(em)-1)):
        raise ValueError("Ex、Em 轴必须分别严格单调，不能有重复或乱序波长")
    if not (ex[0] < 250 < ex[-1] and em[0] < 330 < 380 < em[-1]):
        raise ValueError("五区积分需要 Ex 跨过 250 nm，Em 同时跨过 330 和 380 nm")
    return ex, em, x


def read_matrix(path: Path) -> list[list[float]]:
    """Compatibility helper for the bundled example plot."""
    return read_eem(path)[2]


def line_fill(axis: list[float], values: list[float], keep: list[bool]) -> list[float]:
    known_axis = [a for a, ok in zip(axis, keep) if ok]
    known_values = [v for v, ok in zip(values, keep) if ok]
    if len(known_axis) < 2:
        raise ValueError("散射窗口外不足两个点，无法插值")
    result = []
    for a in axis:
        right = bisect.bisect_left(known_axis, a)
        if right == 0:
            result.append(known_values[0])
        elif right == len(known_axis):
            result.append(known_values[-1])
        elif known_axis[right] == a:
            result.append(known_values[right])
        else:
            left = right - 1
            fraction = (a - known_axis[left]) / (known_axis[right] - known_axis[left])
            result.append(known_values[left] + fraction * (known_values[right] - known_values[left]))
    return result


def interpolate(x: list[list[float]], ex_axis=EX, em_axis=EM):
    if len(x) != len(em_axis) or any(len(row) != len(ex_axis) for row in x):
        raise ValueError("强度矩阵与 Ex、Em 波长轴的长度不一致")
    first = [[abs(em - ex) <= 12 for ex in ex_axis] for em in em_axis]
    second = [[abs(em - (2 * ex + 4)) <= 20 for ex in ex_axis] for em in em_axis]
    mask = [[first[i][j] or second[i][j] for j in range(len(ex_axis))] for i in range(len(em_axis))]
    saturated = sum(value == 9999.9 for row in x for value in row)
    if any(x[i][j] == 9999.9 and not mask[i][j]
           for i in range(len(em_axis)) for j in range(len(ex_axis))):
        raise ValueError("散射窗口外发现 9999.9 饱和值，未自动填补")
    vertical = [line_fill(em_axis, [x[i][j] for i in range(len(em_axis))],
                          [not mask[i][j] for i in range(len(em_axis))]) for j in range(len(ex_axis))]
    horizontal = [line_fill(ex_axis, x[i], [not m for m in mask[i]]) for i in range(len(em_axis))]
    filled = [[max(0.0, (vertical[j][i] + horizontal[i][j]) / 2) if mask[i][j] else x[i][j]
               for j in range(len(ex_axis))] for i in range(len(em_axis))]
    counts = [sum(map(sum, first)),
              sum(second[i][j] and not first[i][j] for i in range(len(em_axis)) for j in range(len(ex_axis))),
              sum(map(sum, mask)), saturated]
    return filled, counts


REGION_PLANES = [
    [(-1, 0, 250), (0, -1, 330), (-1, 1, 0)],
    [(-1, 0, 250), (0, 1, -330), (0, -1, 380), (-1, 1, 0)],
    [(-1, 0, 250), (0, 1, -380), (-1, 1, 0)],
    [(1, 0, -250), (0, -1, 380), (-1, 1, 0)],
    [(1, 0, -250), (0, 1, -380), (-1, 1, 0)],
]


def clip_polygon(poly, plane):
    """Clip (Ex, Em, intensity) vertices against one wavelength half-plane."""
    if not poly:
        return []
    a, b, c = plane
    result = []
    previous = poly[-1]
    fprevious = a * previous[0] + b * previous[1] + c
    for current in poly:
        fcurrent = a * current[0] + b * current[1] + c
        inside_previous, inside_current = fprevious >= 0, fcurrent >= 0
        if inside_previous != inside_current:
            t = fprevious / (fprevious - fcurrent)
            result.append(tuple(previous[k] + t * (current[k] - previous[k]) for k in range(3)))
        if inside_current:
            result.append(current)
        previous, fprevious = current, fcurrent
    return result


def polygon_moments(poly):
    """Exact area and piecewise-linear intensity volume of a clipped triangle."""
    area, volume = 0.0, 0.0
    for k in range(1, len(poly) - 1):
        p, q, r = poly[0], poly[k], poly[k + 1]
        triangle_area = abs((q[0] - p[0]) * (r[1] - p[1]) -
                            (q[1] - p[1]) * (r[0] - p[0])) / 2
        area += triangle_area
        volume += triangle_area * (p[2] + q[2] + r[2]) / 3
    return area, volume


def generic_region_integrals(x, ex_axis, em_axis):
    """Split cells into triangles so FRI cuts need not coincide with grid points."""
    areas = [[] for _ in range(5)]
    volumes = [[] for _ in range(5)]
    z = [[max(0.0, x[i][j]) if em > ex else 0.0
          for j, ex in enumerate(ex_axis)] for i, em in enumerate(em_axis)]
    for i in range(len(em_axis) - 1):
        for j in range(len(ex_axis) - 1):
            p00 = (ex_axis[j], em_axis[i], z[i][j])
            p10 = (ex_axis[j + 1], em_axis[i], z[i][j + 1])
            p11 = (ex_axis[j + 1], em_axis[i + 1], z[i + 1][j + 1])
            p01 = (ex_axis[j], em_axis[i + 1], z[i + 1][j])
            for triangle in ((p00, p10, p11), (p00, p11, p01)):
                for region, planes in enumerate(REGION_PLANES):
                    poly = triangle
                    for plane in planes:
                        poly = clip_polygon(poly, plane)
                        if len(poly) < 3:
                            break
                    if len(poly) >= 3:
                        area, volume = polygon_moments(poly)
                        if area:
                            areas[region].append(area)
                            volumes[region].append(volume)
    return [math.fsum(v) for v in volumes], [math.fsum(a) for a in areas]


def calculate_fri(x: list[list[float]], ex_axis=EX, em_axis=EM):
    if len(x) != len(em_axis) or any(len(row) != len(ex_axis) for row in x):
        raise ValueError("强度矩阵与 Ex、Em 波长轴的长度不一致")
    if list(ex_axis) == EX and list(em_axis) == EM:
        # Preserve the already published outputs for the original instrument grid.
        z = [[max(0.0, x[i][j]) if em > ex else 0.0 for j, ex in enumerate(EX)]
             for i, em in enumerate(EM)]
        phi = []
        for cols, rows in REGIONS:
            total = 0.0
            for ri, i in enumerate(rows):
                row_weight = 0.5 if ri in (0, len(rows) - 1) else 1.0
                for cj, j in enumerate(cols):
                    col_weight = 0.5 if cj in (0, len(cols) - 1) else 1.0
                    total += z[i][j] * row_weight * col_weight
            phi.append(total * 10 * 2)
        areas = AREAS
    else:
        phi, areas = generic_region_integrals(x, ex_axis, em_axis)
    if any(a <= 0 for a in areas):
        raise ValueError("扫描范围内至少一个 FRI 区域没有面积")
    mf = [sum(areas) / a for a in areas]
    normalized = [p * m for p, m in zip(phi, mf)]
    total = sum(normalized)
    if total <= 0 or not math.isfinite(total):
        raise ValueError("面积归一化积分总量不为正")
    return phi + normalized + [v / total for v in normalized] + [total], areas


def integrate(x: list[list[float]], ex_axis=EX, em_axis=EM) -> list[float]:
    return calculate_fri(x, ex_axis, em_axis)[0]


def save_matrix(path: Path, x: list[list[float]], ex_axis=EX, em_axis=EM):
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append([None] + list(ex_axis))
    for em, row in zip(em_axis, x):
        sheet.append([em] + row)
    book.save(path)


def save_summary(path: Path, results, log, area_rows=None, metric_rows=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "积分结果"
    sheet.append(["文件名"] + [f"Phi_{i}" for i in "I II III IV V".split()] +
                 [f"PhiN_{i}" for i in "I II III IV V".split()] +
                 [f"P_{i}" for i in "I II III IV V".split()] + ["PhiN_total"])
    for row in results:
        sheet.append(row)
    sheet = book.create_sheet("荧光指标")
    sheet.append(METRIC_HEADERS)
    for filename, metrics in metric_rows or []:
        sheet.append([filename] + [metrics[key] for key in METRIC_HEADERS[1:]])
    sheet = book.create_sheet("运行记录")
    sheet.append(["文件名", "状态", "说明", "一阶点数", "二阶点数", "插值点数", "饱和点数"])
    for row in log:
        sheet.append(row)
    sheet = book.create_sheet("方法")
    sheet.append(["项目", "规则"])
    for row in [
        ("输入", "标准 XLSX/CSV：首行 Ex、首列 Em；岛津/PicoQuant CSV；日立 F-4600 Data Points TXT；波长轴可使用不同点数和间隔"),
        ("插值", "一阶 |Em−Ex|≤12 nm；二阶 |Em−(2Ex+4)|≤20 nm；沿 Ex 与 Em 双向线性插值取平均，插值负数截 0"),
        ("积分", "负数及 Em≤Ex 记 0；原仪器网格沿用二维梯形法，其他网格以三角网格线性插值并按边界裁剪积分"),
        ("面积归一化", "PhiN_i=Phi_i×(本文件各区域总面积/本文件区域面积_i)；P_i=PhiN_i/sum(PhiN_i)"),
        ("分区边界", "I: Ex≤250/Em≤330；II: Ex≤250/330–380；III: Ex≤250/Em≥380；IV: Ex≥250/Em≤380；V: Ex≥250/Em≥380"),
        ("荧光指标数据", "在瑞利散射插值后的矩阵计算；负强度记 0；非测量波长在原波长网格内线性插值，不外推"),
        ("BIX", "F(Ex310,Em380)/F(Ex310,Em430)；Huguet et al. 2009"),
        ("FI_470_520", "F(Ex370,Em470)/F(Ex370,Em520)；使用 470/520 nm 版本，不是早期 450/500 nm 版本"),
        ("HIX_Ohno", "Ex254 nm；H=Em435–480 nm 积分，L=Em300–345 nm 积分；HIX=H/(L+H)；Ohno 2002 归一化版本"),
        ("Peak_B、Peak_T", "B=F(Ex275,Em310)，T=F(Ex275,Em340)"),
        ("Peak_A、Peak_M、Peak_C", "固定 Ex 下按 Em 区间取峰值及峰位：A=Ex260/Em380–460；M=Ex312/Em380–420；C=Ex350/Em420–480 nm"),
        ("指标缺值", "所需波长不在扫描范围，或比值分母为 0 时对应单元格留空；具体项见荧光指标的备注列"),
        ("用途", "插值点是估计值；不同扫描范围的 FRI 比例不可直接比较，应先统一波长范围"),
    ]:
        sheet.append(row)
    sheet = book.create_sheet("分区面积")
    sheet.append(["文件名", "区域", "面积_nm2", "MF"])
    if area_rows is None:
        area_rows = [("", AREAS)]
    for filename, areas in area_rows:
        for name, area in zip("I II III IV V".split(), areas):
            sheet.append([filename, name, area, sum(areas) / area])
    book.save(path)


def main(input_dir: Path) -> int:
    if not input_dir.is_dir():
        print(f"输入文件夹不存在：{input_dir}")
        return 2
    files = sorted((p for p in input_dir.iterdir() if p.is_file()
                    and p.suffix.lower() in (".xlsx", ".csv", ".txt")
                    and not p.name.startswith("~$")),
                   key=lambda p: p.name)
    if not files:
        print("当前文件夹顶层没有 .xlsx、.csv 或 .txt 文件")
        return 2
    output = input_dir / "EEM_处理结果"
    matrix_dir = output / "01_插值矩阵"
    matrix_dir.mkdir(parents=True, exist_ok=True)
    results, log, area_rows, metric_rows = [], [], [], []
    for path in files:
        print(f"处理：{path.name}", flush=True)
        try:
            ex_axis, em_axis, x = read_eem(path)
            corrected, counts = interpolate(x, ex_axis, em_axis)
            values, areas = calculate_fri(corrected, ex_axis, em_axis)
            metrics = calculate_metrics(corrected, ex_axis, em_axis)
            output_name = path.name if path.suffix.lower() == ".xlsx" else path.name + ".xlsx"
            save_matrix(matrix_dir / output_name, corrected, ex_axis, em_axis)
            results.append([path.name] + values)
            area_rows.append((path.name, areas))
            metric_rows.append((path.name, metrics))
            log.append([path.name, "成功", ""] + counts)
        except Exception as exc:
            print(f"  失败：{exc}", flush=True)
            log.append([path.name, "失败", str(exc), None, None, None, None])
    save_summary(output / "02_五区积分汇总.xlsx", results, log, area_rows, metric_rows)
    failures = sum(row[1] == "失败" for row in log)
    print(f"完成：{len(results)} 成功，{failures} 失败。\n输出目录：{output}")
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        folder = Path(sys.argv[1] if len(sys.argv) > 1 else Path.cwd()).resolve()
        sys.exit(main(folder))
    except Exception as exc:
        print(f"程序错误：{exc}")
        sys.exit(2)
